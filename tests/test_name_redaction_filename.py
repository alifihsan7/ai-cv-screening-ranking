"""A misleading file name must not decide who the candidate is, or the real name survives redaction."""

from __future__ import annotations

import io

import docx

from src.worker.adapter import extract_candidate_profile
from src.worker.redaction import extract_identity_and_redact, redact_text
from src.worker.sanitizer import sanitize_file


def build_cv(name: str = "Arief Nugraha") -> bytes:
    document = docx.Document()
    document.add_paragraph(name)
    document.add_paragraph("Head of Platform Engineering")
    document.add_paragraph(f"Jakarta · 0812-0000-2999 · {name.lower().replace(' ', '.')}@example.com")
    document.add_paragraph("Pengalaman Kerja")
    document.add_paragraph("Head of Platform Engineering — Langit Digital, Jakarta (Apr 2020 – Sekarang)")
    document.add_paragraph("• Memimpin tim platform dan mengelola Kubernetes di AWS.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def redacted_text_for(filename: str) -> tuple[str, dict]:
    content = build_cv()
    profile = extract_candidate_profile(content, filename)
    identity, matching, _ = extract_identity_and_redact(profile)
    text = redact_text(
        sanitize_file(content, filename).extracted_text,
        identity["full_name"],
        [identity["email"]] if identity.get("email") else [],
        [identity["phone"]] if identity.get("phone") else [],
        links=identity.get("links") or [],
    )
    return text, matching


def test_descriptive_file_name_does_not_leave_the_real_name_in_the_redacted_text():
    text, _ = redacted_text_for("K09_platform_head_docx_table2_en.docx")

    assert "arief" not in text.lower()
    assert "nugraha" not in text.lower()


def test_descriptive_file_name_does_not_redact_job_title_words():
    text, matching = redacted_text_for("K09_platform_head_docx_table2_en.docx")

    assert "Head of Platform Engineering" in text
    assert [e["title"] for e in matching["experiences"]] == ["Head of Platform Engineering"]


def test_a_plain_cv_file_name_behaves_the_same():
    text, _ = redacted_text_for("CV_Arief_Nugraha.docx")

    assert "arief" not in text.lower()
    assert "Head of Platform Engineering" in text
