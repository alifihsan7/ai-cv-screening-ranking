"""DOCX CVs keep their tables: job title, company, and dates are often a table row."""

from __future__ import annotations

import io
from pathlib import Path

import docx
import pytest

from src.worker.adapter import extract_candidate_profile
from src.worker.docx_text import extract_docx_text

GOLDEN_CORE = Path("docs/spec/golden_dataset/cvs/core")


def build_docx(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def cv_with_job_table(document: docx.document.Document) -> None:
    document.add_paragraph("Budi Santoso")
    document.add_paragraph("Pengalaman Kerja")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Backend Developer | PT Arunika Digital"
    table.cell(0, 1).text = "Juni 2022 – Sekarang"
    document.add_paragraph("Membangun microservices dengan Python dan PostgreSQL.")
    document.add_paragraph("Pendidikan")
    document.add_paragraph("S1 Teknik Informatika (2016 – 2020)")


def test_text_follows_document_order_with_one_line_per_table_cell():
    text = extract_docx_text(build_docx(cv_with_job_table))

    assert text.splitlines() == [
        "Budi Santoso",
        "Pengalaman Kerja",
        "Backend Developer | PT Arunika Digital",
        "Juni 2022 – Sekarang",
        "Membangun microservices dengan Python dan PostgreSQL.",
        "Pendidikan",
        "S1 Teknik Informatika (2016 – 2020)",
    ]


def test_merged_cells_are_emitted_once():
    def build(document):
        table = document.add_table(rows=1, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Satu sel gabungan"

    assert extract_docx_text(build_docx(build)).splitlines() == ["Satu sel gabungan"]


def test_nested_tables_are_read():
    def build(document):
        outer = document.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "Programmer | CV Mandiri"
        inner.cell(0, 1).text = "2020 – 2022"

    assert extract_docx_text(build_docx(build)).splitlines() == ["Programmer | CV Mandiri", "2020 – 2022"]


def test_job_listed_in_a_table_row_is_parsed_as_experience():
    profile = extract_candidate_profile(build_docx(cv_with_job_table), "cv.docx")

    assert [e["title"] for e in profile["experiences"]] == ["Backend Developer"]
    assert profile["experiences"][0]["start"]["value"] == "2022-06"
    assert "EXPERIENCE_NOT_PARSED" not in profile["quality"]["warnings"]


@pytest.mark.skipif(not GOLDEN_CORE.exists(), reason="golden CVs not available locally")
@pytest.mark.parametrize(
    ("filename", "expected_jobs"),
    [("S5_backend_healthtech_id.docx", 2), ("M4_qa_automation_en.docx", 2), ("I2_accountant_id.docx", 1)],
)
def test_golden_docx_cvs_keep_their_jobs(filename, expected_jobs):
    profile = extract_candidate_profile((GOLDEN_CORE / filename).read_bytes(), filename)

    assert len(profile["experiences"]) == expected_jobs
    assert "EXPERIENCE_NOT_PARSED" not in profile["quality"]["warnings"]


def test_experience_section_that_yields_nothing_is_flagged():
    def build(document):
        document.add_paragraph("Budi Santoso")
        document.add_paragraph("Pengalaman Kerja")
        document.add_paragraph("Pernah bekerja di beberapa perusahaan teknologi tanpa rincian waktu.")

    profile = extract_candidate_profile(build_docx(build), "cv.docx")

    assert profile["experiences"] == []
    assert "EXPERIENCE_NOT_PARSED" in profile["quality"]["warnings"]


def test_cv_without_an_experience_section_is_not_flagged():
    def build(document):
        document.add_paragraph("Budi Santoso")
        document.add_paragraph("Pendidikan")
        document.add_paragraph("S1 Teknik Informatika (2020 – 2024)")

    profile = extract_candidate_profile(build_docx(build), "cv.docx")

    assert profile["experiences"] == []
    assert "EXPERIENCE_NOT_PARSED" not in profile["quality"]["warnings"]


def test_unreadable_docx_is_reported_instead_of_silently_empty():
    profile = extract_candidate_profile(b"this is not a zip archive", "cv.docx")

    assert "DOCX_TEXT_EXTRACTION_FAILED" in profile["quality"]["warnings"]
