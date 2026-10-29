#!/usr/bin/env python3
"""Builds the hold-out set: 12 new fictional CVs in layouts the golden set does not cover, with labels.

Why it exists: every number measured on the golden set comes from 15 CVs that were examined many times while the
system was being fixed. This set is for measuring only. Do NOT tune rules, prompts, thresholds, or settings against
it; if a result disagrees with a label, review the case (the system or the label may be wrong) instead of fitting.

Labels were written from docs/spec/02_SCORING_RUBRIC_SPEC.md (section 4) BEFORE the system was ever run on these
CVs, by the same assistant that worked on the system, so they are not an independent referee like the golden set.
A human recruiter should review them; the `note` of each CV records every non-obvious decision.

Layouts: DOCX with paragraphs only, DOCX with a three-column job table, PDF with prose instead of bullets, and the
golden single-column and two-column PDFs with other date formats.

Usage:  python scripts/build_holdout_set.py     # writes docs/eval/holdout_set/{cvs,labels,README.md}
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
GOLDEN = REPO_ROOT / "docs/spec/golden_dataset"
sys.path.insert(0, str(GOLDEN / "generator"))

import generate_cvs as gen  # noqa: E402
from cv_data import AS_OF, HEADINGS  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from reference_scorer import rank, score  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

S, P, N = "STRONG", "PARTIAL", "NOT_DEMONSTRATED"


def ident(name: str, n: int, city: str, linkedin: str | None = None) -> dict:
    slug = name.lower().replace(" ", ".")
    return {
        "name": name,
        "email": f"{slug}@example.com",
        "phone": f"0812-0000-{2000 + n}",
        "location": city,
        "linkedin": linkedin or f"linkedin.com/in/{slug.replace('.', '-')}-fiktif",
    }


def job(title, org, location, dates, kind, bullets):
    return {"title": title, "org": org, "location": location, "dates": dates, "type": kind, "bullets": bullets}


CVS: list[dict] = [
    {
        "id": "H01", "group": "strong", "format": "docx_para", "lang": "id", "grade": 3,
        "file": "H01_backend_docx_paragraphs_id.docx",
        "identity": ident("Dewi Anggraini", 1, "Bandung"),
        "headline": "Backend Developer",
        "summary": "Backend developer dengan fokus pada layanan penagihan dan pembayaran berbasis Python.",
        "experience": [
            job("Backend Developer", "PT Lintas Data Nusantara", "Bandung", "Januari 2022 – Sekarang", "full_time", [
                "Membangun REST API untuk modul penagihan dengan Python (Django) dan Postgres.",
                "Mengelola deployment layanan di Kubernetes (k8s) pada Amazon Web Services.",
                "Menjalankan layanan dalam container Docker pada lingkungan produksi.",
                "Mempercepat proses penagihan bulanan dari 6 jam menjadi 40 menit.",
            ]),
            job("Junior Programmer", "CV Mitra Solusi", "Surabaya", "Maret 2020 – Desember 2021", "full_time", [
                "Mengembangkan aplikasi web internal dengan Django.",
            ]),
        ],
        "education": [{"degree": "S1 Teknik Informatika", "institution": "Universitas Cendana Raya", "dates": "2015 – 2019"}],
        "certifications": [{"name": "AWS Certified Developer – Associate", "issuer": "Amazon Web Services", "date": "Maret 2023"}],
        "skills": ["Python", "Django", "Postgres", "Docker", "Kubernetes", "AWS", "Git"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [S], "R9": [S], "R10": [N]},
            "backend_months": 80,
            "note": "Alias di teks: 'Postgres' (R4), 'k8s' dan 'Amazon Web Services' (R6). 58 + 22 bulan tanpa tumpang tindih.",
        },
    },
    {
        "id": "H02", "group": "strong", "format": "docx_table3", "lang": "en", "grade": 3,
        "file": "H02_tech_lead_docx_table_en.docx",
        "identity": ident("Adrian Kusuma", 2, "Jakarta"),
        "headline": "Tech Lead, Backend",
        "summary": "Tech lead for commerce platforms built on Python microservices and PostgreSQL.",
        "experience": [
            job("Tech Lead, Backend", "Orbit Commerce", "Jakarta", "Feb 2021 – Present", "full_time", [
                "Led a team of 5 backend engineers building the order and payments platform.",
                "Designed Python (FastAPI) microservices on Google Cloud Platform (GCP) using Cloud Run.",
                "Cut p95 checkout latency from 1.4 s to 320 ms by tuning PostgreSQL queries and adding Redis caching.",
                "Containerized services with Docker.",
            ]),
            job("Backend Engineer", "Pasar Digital", "Jakarta", "Jun 2018 – Jan 2021", "full_time", [
                "Built REST APIs in Python (Django REST Framework) backed by PostgreSQL.",
            ]),
        ],
        "education": [{"degree": "Bachelor of Computer Science", "institution": "Universitas Garuda Timur", "dates": "2014 – 2018"}],
        "certifications": [],
        "skills": ["Python", "FastAPI", "Django", "PostgreSQL", "Redis", "Docker", "GCP"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [N], "R9": [S], "R10": [S]},
            "backend_months": 101,
            "note": "Job header and dates sit in a three-column table. 69 + 32 months. Title 'Tech Lead' plus 'led a team of 5'.",
        },
    },
    {
        "id": "H03", "group": "marginal", "format": "pdf_prose", "lang": "id", "grade": 1,
        "file": "H03_web_developer_prose_id.pdf",
        "heading_overrides": {"experience": "Riwayat Karir"},
        "identity": ident("Rina Hartati", 3, "Yogyakarta"),
        "headline": "Web Developer",
        "summary": "Web developer yang membangun situs e-commerce dan layanan pendukungnya.",
        "experience": [
            job("Web Developer", "Toko Kreatif Digital", "Yogyakarta", "Maret 2024 – Sekarang", "full_time", [
                "Mengembangkan situs e-commerce dengan Django dan MySQL.",
                "Membangun REST API untuk aplikasi mobile toko.",
                "Mengelola basis data MySQL untuk katalog produk.",
            ]),
        ],
        "projects": [job("Proyek Bootcamp: Aplikasi Kasir", "Bootcamp Kode Nusantara", None, "2023", None, [
            "Dibuat dengan Flask dan Docker.",
        ])],
        "education": [{"degree": "S1 Sistem Informasi", "institution": "Universitas Mataram Jaya", "dates": "2017 – 2021"}],
        "certifications": [],
        "skills": ["Python", "Django", "Flask", "MySQL", "Docker"],
        "languages": ["Indonesia"],
        "expected": {
            "states": {"R1": [P], "R2": [S], "R3": [S], "R4": [N], "R5": [P], "R6": [N], "R7": [S], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 32,
            "note": "Heading 'Riwayat Karir' sengaja tidak umum. MySQL bukan alternatif PostgreSQL (spec 4.2). 32 bulan < 36.",
        },
    },
    {
        "id": "H04", "group": "strong", "format": "pdf_single", "lang": "en", "grade": 2,
        "file": "H04_fullstack_numeric_dates_en.pdf",
        "identity": ident("Kevin Santoso", 4, "Surabaya"),
        "headline": "Full Stack Developer",
        "summary": "Full stack developer shipping Python services and React dashboards.",
        "experience": [
            job("Full Stack Developer", "Beacon Retail", "Surabaya", "03/2021 - 02/2024", "full_time", [
                "Built a Python (Flask) REST API and a React admin dashboard for inventory.",
                "Stored data in PostgreSQL and automated deployments with Docker.",
            ]),
            job("Frontend Developer", "Studio Piksel", "Surabaya", "07/2019 - 02/2021", "full_time", [
                "Implemented responsive interfaces with React.",
            ]),
        ],
        "education": [{"degree": "S1 Teknik Informatika", "institution": "Institut Teknologi Arjuna", "dates": "2015 – 2019"}],
        "certifications": [],
        "skills": ["JavaScript", "React", "Python", "Flask", "PostgreSQL", "Docker"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [N], "R7": [S], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 36,
            "note": "Tanggal MM/YYYY. Fullstack diterima; 03/2021-02/2024 = tepat 36 bulan = batas STRONG. Frontend tidak dihitung.",
        },
    },
    {
        "id": "H05", "group": "irrelevant", "format": "pdf_two_col", "lang": "id", "grade": 0,
        "file": "H05_frontend_two_column_id.pdf",
        "identity": ident("Siska Permata", 5, "Malang"),
        "headline": "Frontend Developer",
        "summary": "Frontend developer yang fokus pada antarmuka dashboard dan performa halaman.",
        "experience": [
            job("Frontend Developer", "PT Antar Muka Kreatif", "Malang", "Mei 2021 – Sekarang", "full_time", [
                "Membangun antarmuka dashboard dengan Vue.js dan TypeScript.",
                "Mengintegrasikan antarmuka dengan REST API yang disediakan tim backend.",
                "Mempercepat waktu muat halaman utama sebesar 35 persen.",
            ]),
        ],
        "education": [{"degree": "S1 Teknik Informatika", "institution": "Universitas Kawi Utama", "dates": "2016 – 2020"}],
        "certifications": [],
        "skills": ["JavaScript", "TypeScript", "Vue.js", "HTML", "CSS"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [N], "R2": [P], "R3": [N], "R4": [N], "R5": [N], "R6": [N], "R7": [S], "R8": [N], "R9": [P], "R10": [N]},
            "backend_months": 0,
            "note": "R2: hanya menyebut REST API tanpa membangun = PARTIAL (spec 4.2). R9: angka di pekerjaan tapi domain frontend = PARTIAL.",
        },
    },
    {
        "id": "H06", "group": "marginal", "format": "docx_para", "lang": "en", "grade": 1,
        "file": "H06_data_engineer_docx_en.docx",
        "identity": ident("Tomas Gunawan", 6, "Jakarta"),
        "headline": "Data Engineer",
        "summary": "Data engineer building batch pipelines and warehouses on Google Cloud.",
        "experience": [
            job("Data Engineer", "Insight Labs", "Jakarta", "Sep 2020 – Present", "full_time", [
                "Built batch data pipelines in Python and Airflow on Google Cloud Platform.",
                "Maintained PostgreSQL and BigQuery warehouses for the analytics team.",
                "Reduced nightly pipeline runtime from 5 hours to 90 minutes.",
            ]),
        ],
        "education": [{"degree": "S1 Statistika", "institution": "Universitas Bumi Raya", "dates": "2014 – 2018"}],
        "certifications": [],
        "skills": ["Python", "Airflow", "SQL", "PostgreSQL", "BigQuery", "GCP"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [N], "R2": [N], "R3": [S], "R4": [S], "R5": [N], "R6": [S], "R7": [P], "R8": [N], "R9": [P], "R10": [N]},
            "backend_months": 0,
            "note": "Role family data tidak diterima untuk R1. Statistika = partial_fields. R9: angka di pekerjaan, domain data = PARTIAL.",
        },
    },
    {
        "id": "H07", "group": "marginal", "format": "pdf_prose", "lang": "en", "grade": 1,
        "file": "H07_devops_prose_en.pdf",
        "heading_overrides": {"experience": "Career History"},
        "identity": ident("Yusuf Ramadhan", 7, "Jakarta"),
        "headline": "DevOps Engineer",
        "summary": "DevOps engineer running Kubernetes platforms on public cloud.",
        "experience": [
            job("DevOps Engineer", "CloudNine Systems", "Jakarta", "Jan 2019 – Present", "full_time", [
                "Operated Kubernetes clusters on Amazon Web Services for production workloads.",
                "Built CI/CD pipelines with GitHub Actions.",
                "Containerized legacy services with Docker.",
            ]),
        ],
        "education": [{"degree": "D3 Teknik Komputer", "institution": "Politeknik Cakra Bangsa", "dates": "2013 – 2016"}],
        "certifications": [{"name": "Google Professional Cloud Architect", "issuer": "Google Cloud", "date": "2022"}],
        "skills": ["Kubernetes", "Docker", "AWS", "Terraform", "Linux", "Bash"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [N], "R2": [N], "R3": [N], "R4": [N], "R5": [S], "R6": [S], "R7": [P], "R8": [S], "R9": [N], "R10": [N]},
            "backend_months": 0,
            "note": "Role family devops tidak diterima. D3 di bawah S1 tapi jurusan cocok = PARTIAL (spec 4.2). Sertifikat Google Cloud = STRONG.",
        },
    },
    {
        "id": "H08", "group": "marginal", "format": "docx_table3", "lang": "id", "grade": 1,
        "file": "H08_fresh_graduate_docx_table_id.docx",
        "identity": ident("Fajar Nugroho", 8, "Bekasi"),
        "headline": "Fresh Graduate – Backend Developer",
        "summary": "Lulusan baru Sistem Informasi dengan pengalaman magang backend dan proyek tugas akhir.",
        "experience": [
            job("Backend Developer Intern", "PT Teknologi Maju Bersama", "Bekasi", "Februari 2026 – Juli 2026", "internship", [
                "Membuat endpoint REST API untuk modul absensi dengan Python (FastAPI).",
                "Menulis query PostgreSQL dan memperbaiki bug.",
            ]),
        ],
        "projects": [job("Tugas Akhir: Aplikasi Reservasi Lapangan", "Universitas Nusa Persada", None, "2025", None, [
            "Dibangun dengan FastAPI, PostgreSQL, dan Docker; dipakai 150 pengguna.",
        ])],
        "organizations": [job("Wakil Ketua UKM Pemrograman", "Universitas Nusa Persada", None, "2023 – 2024", None, [
            "Memimpin 12 anggota divisi pembelajaran.",
        ])],
        "education": [{"degree": "S1 Sistem Informasi", "institution": "Universitas Nusa Persada", "dates": "2022 – 2026"}],
        "certifications": [],
        "skills": ["Python", "FastAPI", "PostgreSQL", "Docker", "Git"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [P], "R2": [P], "R3": [S], "R4": [S], "R5": [P], "R6": [N], "R7": [S], "R8": [N], "R9": [P], "R10": [P]},
            "backend_months": 3,
            "note": "Magang 6 bulan x 50% = 3. R2: hanya magang = PARTIAL. R3/R4: pekerjaan bertanggal termasuk magang = STRONG. R5 hanya di proyek.",
        },
    },
    {
        "id": "H09", "group": "irrelevant", "format": "pdf_single", "lang": "id", "grade": 0,
        "file": "H09_career_switcher_id.pdf",
        "identity": ident("Bambang Setiawan", 9, "Cirebon"),
        "headline": "Staf Akuntansi, lulusan bootcamp backend",
        "summary": "Staf akuntansi yang beralih karier ke pengembangan backend lewat bootcamp.",
        "experience": [
            job("Staf Akuntansi", "PT Sumber Rejeki", "Cirebon", "April 2018 – Februari 2025", "full_time", [
                "Menyusun laporan keuangan bulanan dan rekonsiliasi bank.",
                "Mengelola pelaporan pajak PPh 21 dan PPN.",
            ]),
        ],
        "projects": [job("Proyek Bootcamp: Aplikasi Toko Online", "Bootcamp Kode Nusantara", None, "2025", None, [
            "Membangun REST API dengan Python (Flask) dan PostgreSQL.",
        ])],
        "education": [{"degree": "S1 Akuntansi", "institution": "Universitas Dharma Bakti", "dates": "2012 – 2016"}],
        "certifications": [{"name": "Sertifikat Kelulusan Bootcamp Backend", "issuer": "Bootcamp Kode Nusantara", "date": "2025"}],
        "skills": ["Python", "Flask", "PostgreSQL", "Git"],
        "languages": ["Indonesia"],
        "expected": {
            "states": {"R1": [N], "R2": [P], "R3": [P], "R4": [P], "R5": [N], "R6": [N], "R7": [N], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 0,
            "note": "Semua bukti teknis hanya di proyek (PARTIAL). Akuntansi bukan jurusan komputasi. Sertifikat bootcamp bukan sertifikasi cloud.",
        },
    },
    {
        "id": "H10", "group": "strong", "format": "pdf_two_col", "lang": "en", "grade": 3,
        "file": "H10_engineering_manager_two_column_en.pdf",
        "identity": ident("Nadia Prasetyo", 10, "Jakarta"),
        "headline": "Engineering Manager",
        "summary": "Engineering manager with a backend background in ride-matching systems.",
        "experience": [
            job("Engineering Manager", "Zenith Mobility", "Jakarta", "Jan 2023 – Present", "full_time", [
                "Manage three teams (14 engineers) delivering ride-matching services in Go and Python.",
                "Set the architecture direction for services running on Kubernetes (GKE).",
            ]),
            job("Senior Backend Engineer", "Zenith Mobility", "Jakarta", "Aug 2018 – Dec 2022", "full_time", [
                "Designed and built Go and Python microservices handling 2 million ride requests per day.",
                "Introduced PostgreSQL read replicas that cut report query time by 70%.",
            ]),
        ],
        "education": [{"degree": "Master of Computer Science", "institution": "Institut Teknologi Nusantara", "dates": "2016 – 2018"}],
        "certifications": [{"name": "Azure Developer Associate", "issuer": "Microsoft", "date": "2021 (expired 2023)"}],
        "skills": ["Go", "Python", "PostgreSQL", "Kubernetes", "GCP", "Docker"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [P], "R6": [S], "R7": [S], "R8": [P], "R9": [S], "R10": [S]},
            "backend_months": 53,
            "note": "53 bulan sebagai Senior Backend Engineer (99 jika masa manajer dihitung; STRONG di keduanya). Docker hanya di daftar skill = PARTIAL. Sertifikat kedaluwarsa tertulis = PARTIAL.",
        },
    },
    {
        "id": "H11", "group": "marginal", "format": "docx_para", "lang": "id", "grade": 1,
        "file": "H11_qa_lead_docx_id.docx",
        "identity": ident("Ayu Lestari", 11, "Jakarta"),
        "headline": "QA Automation Lead",
        "summary": "QA automation lead yang memimpin tim pengujian layanan keuangan.",
        "experience": [
            job("QA Automation Lead", "PT Quality First", "Jakarta", "Juni 2020 – Sekarang", "full_time", [
                "Memimpin tim QA 4 orang dan menyusun strategi pengujian otomatis.",
                "Membuat suite pengujian API dengan Python (pytest) untuk 90 endpoint REST.",
                "Menjalankan lingkungan uji terisolasi dengan Docker.",
                "Mengurangi waktu pengujian regresi dari 3 hari menjadi 5 jam.",
            ]),
        ],
        "education": [{"degree": "S1 Teknik Informatika", "institution": "Universitas Wijaya Mulia", "dates": "2014 – 2018"}],
        "certifications": [],
        "skills": ["Python", "pytest", "Selenium", "Docker", "PostgreSQL"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [N], "R2": [N, P], "R3": [S], "R4": [P], "R5": [S], "R6": [N], "R7": [S], "R8": [N], "R9": [P], "R10": [S, P]},
            "backend_months": 0,
            "note": "Mengikuti konvensi golden M4: menguji API bukan membangun (NOT_DEMONSTRATED atau PARTIAL); tim QA untuk lowongan backend boleh STRONG atau PARTIAL. PostgreSQL hanya di daftar skill.",
        },
    },
    {
        "id": "H12", "group": "strong", "format": "pdf_prose", "lang": "id", "grade": 2,
        "file": "H12_java_backend_prose_id.pdf",
        "heading_overrides": {"experience": "Pengalaman Profesional"},
        "identity": ident("Wahyu Hidayat", 12, "Jakarta"),
        "headline": "Java Backend Developer",
        "summary": "Backend developer berpengalaman di sistem transaksi perbankan berbasis Java.",
        "experience": [
            job("Java Backend Developer", "PT Bank Digital Sentosa", "Jakarta", "Agustus 2018 – Sekarang", "full_time", [
                "Membangun layanan microservices berbasis Java (Spring Boot) untuk transaksi perbankan.",
                "Mengelola basis data PostgreSQL dan mengoptimalkan query yang lambat.",
                "Men-deploy layanan dengan Docker dan Kubernetes di Microsoft Azure.",
                "Menurunkan waktu proses transaksi batch dari 45 menit menjadi 12 menit.",
            ]),
        ],
        "education": [{"degree": "S1 Ilmu Komputer", "institution": "Universitas Sriwijaya Baru", "dates": "2013 – 2017"}],
        "certifications": [],
        "skills": ["Java", "Spring Boot", "PostgreSQL", "Docker", "Kubernetes", "Azure"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [N], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [N], "R9": [S], "R10": [N]},
            "backend_months": 99,
            "note": "Tanpa Python (Java bukan alternatif, spec 4.2 butir 5). Azure di pekerjaan = STRONG untuk R6.",
        },
    },
]


# ----------------------------------------------------------------- renderers
def _text(cv: dict, key: str) -> str:
    return cv.get("heading_overrides", {}).get(key, HEADINGS[cv["lang"]][key])


def _header(doc: Document, cv: dict) -> None:
    doc.add_paragraph().add_run(cv["identity"]["name"]).bold = True
    doc.add_paragraph(cv["headline"])
    doc.add_paragraph(gen.contact_line(cv))
    doc.add_paragraph().add_run(_text(cv, "summary")).bold = True
    doc.add_paragraph(cv["summary"])


def _tail(doc: Document, cv: dict) -> None:
    for key in ("projects", "organizations"):
        for entry in cv.get(key, []):
            if entry is cv[key][0]:
                doc.add_paragraph().add_run(_text(cv, key)).bold = True
            doc.add_paragraph().add_run(f"{entry['title']} — {entry['org']} ({entry['dates']})").bold = True
            for b in entry["bullets"]:
                doc.add_paragraph(f"• {b}")
    doc.add_paragraph().add_run(_text(cv, "education")).bold = True
    for ed in cv["education"]:
        doc.add_paragraph(f"{ed['degree']} – {ed['institution']} ({ed['dates']})")
    if cv.get("certifications"):
        doc.add_paragraph().add_run(_text(cv, "certifications")).bold = True
        for c in cv["certifications"]:
            doc.add_paragraph(f"{c['name']} — {c['issuer']} ({c['date']})")
    doc.add_paragraph().add_run(_text(cv, "skills")).bold = True
    doc.add_paragraph(", ".join(cv["skills"]))
    doc.add_paragraph().add_run(_text(cv, "languages")).bold = True
    doc.add_paragraph(", ".join(cv["languages"]))


def _new_doc() -> Document:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    return doc


def render_docx_para(cv: dict, path: Path) -> None:
    """Paragraphs only: 'Title — Org, City (dates)' as one bold line, bullets as paragraphs."""
    doc = _new_doc()
    _header(doc, cv)
    doc.add_paragraph().add_run(_text(cv, "experience")).bold = True
    for e in cv["experience"]:
        doc.add_paragraph().add_run(f"{e['title']} — {e['org']}, {e['location']} ({e['dates']})").bold = True
        for b in e["bullets"]:
            doc.add_paragraph(f"• {b}")
    _tail(doc, cv)
    doc.save(path)


def render_docx_table3(cv: dict, path: Path) -> None:
    """One row per job with three cells: dates | title and organization | bullets as paragraphs."""
    doc = _new_doc()
    _header(doc, cv)
    doc.add_paragraph().add_run(_text(cv, "experience")).bold = True
    for e in cv["experience"]:
        table = doc.add_table(rows=1, cols=3)
        dates, role, detail = table.rows[0].cells
        dates.text = e["dates"]
        role.text = e["title"]
        role.add_paragraph(f"{e['org']}, {e['location']}")
        detail.text = f"• {e['bullets'][0]}"
        for b in e["bullets"][1:]:
            detail.add_paragraph(f"• {b}")
        doc.add_paragraph()
    _tail(doc, cv)
    doc.save(path)


def render_pdf_prose(cv: dict, path: Path) -> None:
    """Experience as running prose: a header line with inline dates, then the bullets as one paragraph."""
    story = [
        Paragraph(cv["identity"]["name"].upper(), gen.ST["name"]),
        Paragraph(cv["headline"], gen.ST["headline"]),
        Spacer(1, 3),
        Paragraph(gen.contact_line(cv), gen.ST["contact"]),
        Paragraph(_text(cv, "summary"), gen.ST["h"]),
        Paragraph(cv["summary"], gen.ST["body"]),
        Paragraph(_text(cv, "experience"), gen.ST["h"]),
    ]
    for e in cv["experience"]:
        story.append(Paragraph(f"<b>{e['title']}</b>, {e['org']}, {e['location']}. {e['dates']}.", gen.ST["body"]))
        story.append(Paragraph(" ".join(e["bullets"]), gen.ST["body"]))
        story.append(Spacer(1, 4))
    for key in ("projects", "organizations"):
        entries = cv.get(key, [])
        if entries:
            story.append(Paragraph(_text(cv, key), gen.ST["h"]))
            for e in entries:
                story.append(Paragraph(f"<b>{e['title']}</b>, {e['org']}, {e['dates']}. {' '.join(e['bullets'])}", gen.ST["body"]))
    story.append(Paragraph(_text(cv, "education"), gen.ST["h"]))
    for ed in cv["education"]:
        story.append(Paragraph(f"{ed['degree']} – {ed['institution']} ({ed['dates']})", gen.ST["body"]))
    if cv.get("certifications"):
        story.append(Paragraph(_text(cv, "certifications"), gen.ST["h"]))
        for c in cv["certifications"]:
            story.append(Paragraph(f"{c['name']} — {c['issuer']} ({c['date']})", gen.ST["body"]))
    story += [Paragraph(_text(cv, "skills"), gen.ST["h"]), Paragraph(", ".join(cv["skills"]), gen.ST["body"])]
    story += [Paragraph(_text(cv, "languages"), gen.ST["h"]), Paragraph(", ".join(cv["languages"]), gen.ST["body"])]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm).build(story)


RENDER = {
    "docx_para": render_docx_para,
    "docx_table3": render_docx_table3,
    "pdf_prose": render_pdf_prose,
    "pdf_single": gen.render_pdf_single,
    "pdf_two_col": gen.render_pdf_two_col,
}


def build_labels() -> dict:
    jd = json.loads((GOLDEN / "job_descriptions/JD-01_requirements.approved.json").read_text())
    candidates = []
    scored = []
    for cv in CVS:
        states = cv["expected"]["states"]
        candidates.append({
            "cv_id": cv["id"],
            "file": f"cvs/core/{cv['file']}",
            "group": cv["group"],
            "format": cv["format"],
            "language": cv["lang"],
            "relevance_grade": cv["grade"],
            "expected_states": states,
            "expected_backend_months": cv["expected"]["backend_months"],
            "note": cv["expected"]["note"],
        })
        first = {k: v[0] for k, v in states.items()}
        scored.append({
            "candidate_id": cv["id"],
            "result": score(first, jd),
            "strong_count": sum(v == S for v in first.values()),
        })
    ranked = rank(scored)
    return {
        "job_id": "JD-01",
        "as_of": AS_OF,
        "requirement_set": "docs/spec/golden_dataset/job_descriptions/JD-01_requirements.approved.json",
        "authoring": "Labels derived from docs/spec/02_SCORING_RUBRIC_SPEC.md section 4 before the system was run on "
                     "these CVs. Written by the assistant that worked on the system; needs human review.",
        "reference_ranking": [r["candidate_id"] for r in ranked],
        "candidates": candidates,
    }


def readme(labels: dict) -> str:
    rows = []
    for c in labels["candidates"]:
        st = " ".join("/".join(s[0] for s in c["expected_states"][f"R{i}"]) for i in range(1, 11))
        rows.append(f"| {c['cv_id']} | {c['format']} | {c['language']} | {c['relevance_grade']} | `{st}` | {c['note']} |")
    return (
        "# Hold-out set\n\n"
        "Twelve new fictional CVs in layouts the golden dataset does not cover (DOCX with paragraphs only, DOCX with a "
        "three-column job table, PDF with prose instead of bullets, other date formats), scored against JD-01.\n\n"
        "**Measure only. Do not tune rules, prompts, thresholds, or settings against this set.** If the system disagrees "
        "with a label, review the case: either may be wrong. The golden set stays the referee.\n\n"
        "**The labels are not independent.** They were written from the scoring spec (section 4) before the system was "
        "run on these CVs, but by the assistant that worked on the system. A human recruiter should review them.\n\n"
        "Rebuild the CV files (PDF and DOCX are git-ignored) with `python scripts/build_holdout_set.py`, then:\n\n"
        "```\npython scripts/eval_agent_path.py --cvs-dir docs/eval/holdout_set/cvs/core "
        "--labels docs/eval/holdout_set/labels/expected_JD-01.json --output-json out.json\n```\n\n"
        "## Expected states (R1 to R10: S = STRONG, P = PARTIAL, N = NOT_DEMONSTRATED)\n\n"
        "| CV | Layout | Lang | Grade | R1..R10 | Notes |\n|---|---|---|---|---|---|\n" + "\n".join(rows) + "\n"
    )


def main() -> None:
    ap = argparse.ArgumentParser(description="Build the hold-out set")
    ap.add_argument("--out", type=Path, default=REPO_ROOT / "docs/eval/holdout_set")
    args = ap.parse_args()

    core = args.out / "cvs" / "core"
    core.mkdir(parents=True, exist_ok=True)
    for cv in CVS:
        RENDER[cv["format"]](cv, core / cv["file"])
        print("wrote", cv["file"])

    labels = build_labels()
    (args.out / "labels").mkdir(parents=True, exist_ok=True)
    (args.out / "labels" / "expected_JD-01.json").write_text(json.dumps(labels, indent=2, ensure_ascii=False))
    (args.out / "README.md").write_text(readme(labels))
    print(f"\n{len(CVS)} CVs -> {args.out}; reference ranking: {labels['reference_ranking']}")


if __name__ == "__main__":
    main()
