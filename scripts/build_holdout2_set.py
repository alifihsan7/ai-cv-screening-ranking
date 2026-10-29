#!/usr/bin/env python3
"""Builds the second hold-out set: 12 more fictional CVs (K01..K12) with labels.

Same discipline as the first hold-out (scripts/build_holdout_set.py): measure only, never tune rules, prompts,
thresholds, or settings against it. Labels come from docs/spec/02_SCORING_RUBRIC_SPEC.md (section 4), written
BEFORE the system was run on these CVs, by the assistant that worked on the system; a human should review them.

This set is not a random sample. It was designed after the first hold-out was measured, so it deliberately
exercises behaviors already under suspicion:
  - ACHIEVEMENT with no numeric result at all, and with numbers in an unrelated domain (K05, K06, K07, K10, K11),
  - a numeric result written as a word ("tiga kali lipat", K01) and a numeric result only in a project (K02, K08),
  - a "date first, then title" layout (K03). It has a single job, and the known parser defect only shows with two or
    more jobs, so K03 does not actually exercise it.
It also adds layouts the first sets lack: a CV laid out as one big two-column table (K07), jobs in a two-column
table (K01, K09), inline "Title at Org (dates)" headers (K02, K10), and year-only dates (K02).

Usage:  python scripts/build_holdout2_set.py     # writes docs/eval/holdout_set_2/{cvs,labels,README.md}
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_holdout_set as h1  # noqa: E402
from build_holdout_set import GOLDEN, N, P, S, ident, job  # noqa: E402
from cv_data import AS_OF  # noqa: E402
from docx.shared import Pt  # noqa: E402
from reference_scorer import rank, score  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

gen = h1.gen

CVS: list[dict] = [
    {
        "id": "K01", "group": "strong", "format": "docx_table2", "lang": "id", "grade": 3,
        "file": "K01_backend_docx_table2_id.docx",
        "identity": ident("Hendra Wirawan", 21, "Jakarta"),
        "headline": "Backend Engineer",
        "summary": "Backend engineer untuk layanan pembayaran dompet digital.",
        "experience": [
            job("Backend Engineer", "PT Arta Pay", "Jakarta", "Mei 2021 – Sekarang", "full_time", [
                "Membangun REST API untuk dompet digital dengan Python (FastAPI) dan PostgreSQL.",
                "Meningkatkan throughput API pembayaran menjadi tiga kali lipat lewat caching Redis.",
                "Men-deploy layanan ke Amazon EKS dengan Docker.",
            ]),
            job("Software Engineer", "PT Nusa Solusi", "Bogor", "Juli 2019 – April 2021", "full_time", [
                "Mengembangkan layanan backend dengan Django.",
            ]),
        ],
        "education": [{"degree": "S1 Ilmu Komputer", "institution": "Universitas Pelita Harapan Baru", "dates": "2015 – 2019"}],
        "certifications": [{"name": "AWS Certified Solutions Architect – Associate", "issuer": "Amazon Web Services", "date": "2022"}],
        "skills": ["Python", "FastAPI", "Django", "PostgreSQL", "Redis", "Docker", "AWS"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [S], "R9": [S], "R10": [N]},
            "backend_months": 88,
            "note": "Hasil terukur ditulis dengan kata ('tiga kali lipat') = angka hasil di pekerjaan backend. 'Amazon EKS' = Kubernetes di AWS. 66 + 22 bulan.",
        },
    },
    {
        "id": "K02", "group": "strong", "format": "pdf_inline", "lang": "en", "grade": 2,
        "file": "K02_backend_inline_year_only_en.pdf",
        "identity": ident("Oliver Tanuwijaya", 22, "Bandung"),
        "headline": "Backend Developer",
        "summary": "Backend developer on billing services, mostly Python and PostgreSQL.",
        "experience": [
            job("Backend Developer", "Kite Systems", "Bandung", "2022 - Present", "full_time", [
                "Developed REST APIs in Python (Django) for the billing service in production.",
                "Maintained PostgreSQL databases and wrote migrations.",
            ]),
            job("Backend Developer", "Roda Teknologi", "Semarang", "2019 - 2022", "full_time", [
                "Built internal tools in Python (Flask).",
            ]),
        ],
        "projects": [job("Personal project: Expense Tracker API", "GitHub", None, "2024", None, [
            "Handled 5,000 requests per day in a load test.",
        ])],
        "education": [{"degree": "Bachelor of Information Systems", "institution": "Institut Teknologi Merdeka", "dates": "2015 – 2019"}],
        "certifications": [{"name": "Microsoft Azure Fundamentals (AZ-900)", "issuer": "Microsoft", "date": "2023"}],
        "skills": ["Python", "Django", "Flask", "PostgreSQL", "Azure"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [N], "R6": [P], "R7": [S], "R8": [S], "R9": [P], "R10": [N]},
            "backend_months": 88,
            "note": "Tanggal hanya tahun: '2019 - 2022' = Jul 2019-Jun 2022 = 36 bulan, '2022 - Present' = Jul 2022-Okt 2026 = 52 bulan (aturan year-only). Azure hanya di daftar skill = PARTIAL. Angka hanya di proyek = PARTIAL.",
        },
    },
    {
        "id": "K03", "group": "marginal", "format": "docx_dates_first", "lang": "id", "grade": 2,
        "file": "K03_backend_docx_dates_first_id.docx",
        "identity": ident("Citra Maheswari", 23, "Medan"),
        "headline": "Backend Developer",
        "summary": "Backend developer untuk sistem kasir toko ritel.",
        "experience": [
            job("Backend Developer", "PT Sigma Retail", "Medan", "Januari 2024 – Sekarang", "full_time", [
                "Mengembangkan REST API untuk sistem kasir dengan Python (Flask).",
                "Menggunakan PostgreSQL untuk penyimpanan transaksi.",
                "Mengurangi waktu sinkronisasi stok dari 20 menit menjadi 4 menit.",
            ]),
        ],
        "education": [{"degree": "S1 Teknik Informatika", "institution": "Universitas Sumatera Maju", "dates": "2019 – 2023"}],
        "certifications": [],
        "skills": ["Python", "Flask", "PostgreSQL", "Git"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [P], "R2": [S], "R3": [S], "R4": [S], "R5": [N], "R6": [N], "R7": [S], "R8": [N], "R9": [S], "R10": [N]},
            "backend_months": 34,
            "note": "Tata letak 'tanggal dulu, baru judul', tetapi hanya satu pekerjaan, sehingga cacat parser yang diketahui (butuh dua pekerjaan atau lebih) tidak teruji di sini. Jan 2024-Okt 2026 = 34 bulan < 36.",
        },
    },
    {
        "id": "K04", "group": "strong", "format": "pdf_single", "lang": "en", "grade": 2,
        "file": "K04_senior_java_single_en.pdf",
        "identity": ident("Raymond Halim", 24, "Jakarta"),
        "headline": "Senior Software Engineer",
        "summary": "Senior engineer building payment microservices in Java.",
        "experience": [
            job("Senior Software Engineer", "Mantra Bank", "Jakarta", "Mar 2017 – Present", "full_time", [
                "Designed and built payment microservices in Java (Spring Boot) serving 1.5 million transactions per day.",
                "Mentored three junior engineers through weekly code reviews.",
                "Ran services on Kubernetes in Microsoft Azure.",
                "Tuned PostgreSQL indexes for the ledger database.",
            ]),
        ],
        "education": [{"degree": "Bachelor of Software Engineering", "institution": "Institut Teknologi Sakti", "dates": "2012 – 2016"}],
        "certifications": [],
        "skills": ["Java", "Spring Boot", "PostgreSQL", "Kubernetes", "Azure"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [N], "R4": [S], "R5": [N], "R6": [S], "R7": [S], "R8": [N], "R9": [S], "R10": [S]},
            "backend_months": 116,
            "note": "Tanpa Python. Docker tidak disebut. 'Mentored three junior engineers' = mentoring tim teknis = STRONG untuk R10.",
        },
    },
    {
        "id": "K05", "group": "irrelevant", "format": "docx_para", "lang": "id", "grade": 0,
        "file": "K05_digital_marketing_docx_id.docx",
        "identity": ident("Rizal Fahmi", 25, "Bekasi"),
        "headline": "Digital Marketing Specialist",
        "summary": "Spesialis pemasaran digital untuk merek konsumen.",
        "experience": [
            job("Digital Marketing Specialist", "PT Merek Nusantara", "Bekasi", "Maret 2019 – Sekarang", "full_time", [
                "Menjalankan kampanye iklan berbayar di media sosial dan mesin pencari.",
                "Meningkatkan tingkat konversi situs sebesar 28 persen.",
                "Memimpin tim konten beranggotakan 3 orang.",
            ]),
        ],
        "education": [{"degree": "S1 Manajemen", "institution": "Universitas Bisnis Cipta", "dates": "2014 – 2018"}],
        "certifications": [],
        "skills": ["Google Analytics", "Meta Ads", "SEO", "Excel"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [N], "R2": [N], "R3": [N], "R4": [N], "R5": [N], "R6": [N], "R7": [N], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 0,
            "note": "Angka hasil ada tetapi di domain pemasaran (domain sama sekali berbeda) = NOT_DEMONSTRATED. Tim konten non-teknis = NOT_DEMONSTRATED untuk R10.",
        },
    },
    {
        "id": "K06", "group": "marginal", "format": "pdf_prose", "lang": "id", "grade": 1,
        "file": "K06_fullstack_node_prose_id.pdf",
        "identity": ident("Gita Maharani", 26, "Denpasar"),
        "headline": "Full Stack Developer",
        "summary": "Full stack developer untuk aplikasi reservasi berbasis Node.js dan React.",
        "experience": [
            job("Full Stack Developer", "PT Kreasi Aplikasi", "Denpasar", "Oktober 2022 – Sekarang", "full_time", [
                "Membuat REST API dengan Node.js dan PostgreSQL untuk aplikasi reservasi.",
                "Mengembangkan antarmuka dengan React.",
                "Menjalankan layanan dalam container Docker.",
            ]),
        ],
        "education": [{"degree": "D4 Teknik Informatika", "institution": "Politeknik Pesisir Utama", "dates": "2018 – 2022"}],
        "certifications": [],
        "skills": ["JavaScript", "Node.js", "React", "PostgreSQL", "Docker"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [N], "R4": [S], "R5": [S], "R6": [N], "R7": [S], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 49,
            "note": "Fullstack diterima untuk R1 (Okt 2022-Okt 2026 = 49 bulan). D4 setara S1 dengan jurusan utama = STRONG. Tidak ada angka hasil sama sekali.",
        },
    },
    {
        "id": "K07", "group": "marginal", "format": "docx_sidebar", "lang": "en", "grade": 2,
        "file": "K07_backend_docx_sidebar_en.docx",
        "identity": ident("Priya Wijaya", 27, "Jakarta"),
        "headline": "Backend Developer",
        "summary": "Backend developer building seller tooling for an online marketplace.",
        "experience": [
            job("Backend Developer", "Tokoloka", "Jakarta", "Jan 2022 – Present", "full_time", [
                "Built Go and Python (FastAPI) microservices for the seller dashboard.",
                "Deployed services to Google Cloud Platform (GCP) with Docker.",
                "Added load tests to the release pipeline.",
            ]),
        ],
        "education": [{"degree": "S1 Matematika", "institution": "Universitas Angka Utama", "dates": "2013 – 2017"}],
        "certifications": [{"name": "Google Associate Cloud Engineer", "issuer": "Google Cloud", "date": "2023"}],
        "skills": ["Go", "Python", "FastAPI", "Docker", "GCP"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [N], "R5": [S], "R6": [S], "R7": [P], "R8": [S], "R9": [N], "R10": [N]},
            "backend_months": 58,
            "note": "Seluruh CV berupa satu tabel dua kolom (sidebar). Matematika ada di partial_fields = PARTIAL. PostgreSQL tidak disebut. Tidak ada angka hasil.",
        },
    },
    {
        "id": "K08", "group": "marginal", "format": "pdf_two_col", "lang": "id", "grade": 1,
        "file": "K08_active_student_two_column_id.pdf",
        "identity": ident("Dimas Prakoso", 28, "Depok"),
        "headline": "Mahasiswa Teknik Informatika",
        "summary": "Mahasiswa semester 7 yang mencari posisi backend developer.",
        "experience": [],
        "projects": [job("Tugas Kuliah: Aplikasi Pemesanan Tiket", "Universitas Madani Jaya", None, "2025", None, [
            "Membangun REST API pemesanan tiket dengan Python (Flask) dan PostgreSQL; dijalankan dengan Docker dan diuji oleh 80 pengguna.",
        ])],
        "organizations": [job("Anggota Himpunan Mahasiswa Informatika", "Universitas Madani Jaya", None, "2023 – 2024", None, [
            "Membantu penyelenggaraan seminar teknologi tahunan.",
        ])],
        "education": [{"degree": "S1 Teknik Informatika (mahasiswa aktif, semester 7)", "institution": "Universitas Madani Jaya", "dates": "2022 – sekarang"}],
        "certifications": [],
        "skills": ["Python", "Flask", "PostgreSQL", "Docker"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [N], "R2": [P], "R3": [P], "R4": [P], "R5": [P], "R6": [N], "R7": [P], "R8": [N], "R9": [P], "R10": [N]},
            "backend_months": 0,
            "note": "Tanpa pengalaman kerja: semua bukti di proyek (PARTIAL). Mahasiswa aktif belum lulus = maksimal PARTIAL (spec 4.2 EDUCATION). Angka hasil hanya di proyek = PARTIAL.",
        },
    },
    {
        "id": "K09", "group": "marginal", "format": "docx_table2", "lang": "en", "grade": 1,
        "file": "K09_platform_head_docx_table2_en.docx",
        "identity": ident("Arief Nugraha", 29, "Jakarta"),
        "headline": "Head of Platform Engineering",
        "summary": "Platform and reliability engineer leading an infrastructure team.",
        "experience": [
            job("Head of Platform Engineering", "Langit Digital", "Jakarta", "Apr 2020 – Present", "full_time", [
                "Built internal command-line tooling for deployments.",
                "Wrote automation scripts in Python for capacity reports.",
                "Cut p99 API latency by 45% by tuning autoscaling on Kubernetes.",
                "Ran container workloads with Docker on AWS.",
            ]),
        ],
        "education": [{"degree": "S1 Ilmu Komputer", "institution": "Universitas Andalas Baru", "dates": "2011 – 2015"}],
        "certifications": [{"name": "AWS Certified SysOps Administrator – Associate", "issuer": "Amazon Web Services", "date": "2021"}],
        "skills": ["Kubernetes", "Docker", "AWS", "Python", "Terraform"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [N], "R2": [N], "R3": [S], "R4": [N], "R5": [S], "R6": [S], "R7": [S], "R8": [S], "R9": [S], "R10": [S]},
            "backend_months": 0,
            "note": "Role family devops tidak diterima untuk R1. Tooling CLI bukan REST API. Python di pekerjaan = STRONG. 'API latency' = kinerja backend = angka hasil relevan. Judul 'Head of' = STRONG untuk R10.",
        },
    },
    {
        "id": "K10", "group": "strong", "format": "pdf_inline", "lang": "id", "grade": 3,
        "file": "K10_backend_inline_overlap_id.pdf",
        "identity": ident("Yoga Pratama", 30, "Surabaya"),
        "headline": "Backend Developer",
        "summary": "Backend developer di perusahaan logistik, juga mengerjakan proyek lepas.",
        "experience": [
            job("Backend Developer", "PT Cakra Teknologi", "Surabaya", "Februari 2019 – Desember 2023", "full_time", [
                "Mengembangkan microservices pengiriman dengan Java dan Kotlin.",
                "Mengelola basis data PostgreSQL.",
            ]),
            job("Freelance Backend Developer", "Proyek lepas", "Remote", "Januari 2022 – Desember 2022", "freelance", [
                "Membuat API untuk toko online dengan Python (Django).",
            ]),
            job("Backend Developer", "PT Garuda Niaga", "Surabaya", "Januari 2024 – Sekarang", "full_time", [
                "Membangun REST API dengan Python (FastAPI).",
                "Menggunakan Docker dan Kubernetes (k8s) di AWS.",
            ]),
        ],
        "education": [{"degree": "S1 Informatika", "institution": "Institut Sains Teknologi Timur", "dates": "2014 – 2018"}],
        "certifications": [],
        "skills": ["Java", "Kotlin", "Python", "FastAPI", "Django", "PostgreSQL", "Docker", "Kubernetes", "AWS"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 93,
            "note": "Freelance 2022 tumpang tindih penuh dengan pekerjaan utama: tidak dijumlahkan dua kali (59 + 34 = 93). Tidak ada angka hasil sama sekali.",
        },
    },
    {
        "id": "K11", "group": "irrelevant", "format": "docx_para", "lang": "en", "grade": 0,
        "file": "K11_warehouse_supervisor_docx_en.docx",
        "identity": ident("Edo Saputra", 31, "Cikarang"),
        "headline": "Warehouse Operations Supervisor, backend bootcamp graduate",
        "summary": "Operations supervisor moving into backend development after a bootcamp.",
        "experience": [
            job("Warehouse Operations Supervisor", "Gudang Sentosa", "Cikarang", "2018 – 2024", "full_time", [
                "Supervised a team of 8 warehouse staff.",
                "Reduced picking errors by 22 percent.",
            ]),
        ],
        "projects": [job("Bootcamp project: Inventory API", "Bootcamp Kode Nusantara", None, "2024", None, [
            "Built a REST API with Python (FastAPI) and PostgreSQL, packaged with Docker and deployed on AWS.",
        ])],
        "education": [{"degree": "S1 Teknik Industri", "institution": "Universitas Karya Mandiri", "dates": "2013 – 2017"}],
        "certifications": [],
        "skills": ["Python", "FastAPI", "PostgreSQL", "Docker", "AWS"],
        "languages": ["Indonesian", "English"],
        "expected": {
            "states": {"R1": [N], "R2": [P], "R3": [P], "R4": [P], "R5": [P], "R6": [P], "R7": [N], "R8": [N], "R9": [N], "R10": [N]},
            "backend_months": 0,
            "note": "Bukti teknis hanya di proyek bootcamp (PARTIAL). Angka hasil (22 persen) ada tetapi di domain gudang = NOT_DEMONSTRATED. Memimpin tim non-teknis = NOT_DEMONSTRATED untuk R10. Teknik Industri bukan jurusan komputasi.",
        },
    },
    {
        "id": "K12", "group": "strong", "format": "pdf_single", "lang": "id", "grade": 3,
        "file": "K12_team_coordinator_single_id.pdf",
        "heading_overrides": {"experience": "Pengalaman"},
        "identity": ident("Lina Marlina", 32, "Jakarta"),
        "headline": "Koordinator Tim Backend",
        "summary": "Backend developer yang mengoordinasikan tim layanan pinjaman digital.",
        "experience": [
            job("Koordinator Tim Backend", "PT Dana Cepat", "Jakarta", "Maret 2022 – Sekarang", "full_time", [
                "Membangun REST API pinjaman dengan Python (Django) dan PostgreSQL.",
                "Mengoordinasikan 5 engineer dalam tim backend.",
                "Mempercepat pencairan dana 2x lipat dengan antrean asinkron.",
                "Men-deploy layanan ke Google Cloud Platform menggunakan Docker.",
            ]),
        ],
        "education": [{"degree": "S1 Sistem Informasi", "institution": "Universitas Bina Cendekia Baru", "dates": "2016 – 2020"}],
        "certifications": [],
        "skills": ["Python", "Django", "PostgreSQL", "Docker", "GCP"],
        "languages": ["Indonesia", "Inggris"],
        "expected": {
            "states": {"R1": [S], "R2": [S], "R3": [S], "R4": [S], "R5": [S], "R6": [S], "R7": [S], "R8": [N], "R9": [S], "R10": [S]},
            "backend_months": 56,
            "note": "Judul 'Koordinator Tim Backend' = STRONG untuk R10 (spec 4.2 LEADERSHIP). Heading 'Pengalaman' saja.",
        },
    },
]


# ----------------------------------------------------------------- new renderers
def _bold_line(doc, text: str) -> None:
    doc.add_paragraph().add_run(text).bold = True


def render_docx_table2(cv: dict, path: Path) -> None:
    """One row per job with two cells: title, organization and dates on the left; bullets on the right."""
    doc = h1._new_doc()
    h1._header(doc, cv)
    _bold_line(doc, h1._text(cv, "experience"))
    for e in cv["experience"]:
        table = doc.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        left.text = e["title"]
        left.add_paragraph(f"{e['org']}, {e['location']}")
        left.add_paragraph(e["dates"])
        right.text = f"• {e['bullets'][0]}"
        for b in e["bullets"][1:]:
            right.add_paragraph(f"• {b}")
        doc.add_paragraph()
    h1._tail(doc, cv)
    doc.save(path)


def render_docx_dates_first(cv: dict, path: Path) -> None:
    """Paragraphs where the date line comes first, then the title, then the organization, then bullets."""
    doc = h1._new_doc()
    h1._header(doc, cv)
    _bold_line(doc, h1._text(cv, "experience"))
    for e in cv["experience"]:
        doc.add_paragraph(e["dates"])
        _bold_line(doc, e["title"])
        doc.add_paragraph(f"{e['org']}, {e['location']}")
        for b in e["bullets"]:
            doc.add_paragraph(f"• {b}")
    h1._tail(doc, cv)
    doc.save(path)


def render_docx_sidebar(cv: dict, path: Path) -> None:
    """The whole CV below the name is one two-column table: sidebar on the left, main content on the right."""
    doc = h1._new_doc()
    doc.add_paragraph().add_run(cv["identity"]["name"]).bold = True
    doc.add_paragraph(cv["headline"])
    table = doc.add_table(rows=1, cols=2)
    side, main = table.rows[0].cells

    def put(cell, text: str, bold: bool = False) -> None:
        paragraph = cell.paragraphs[0] if not cell.paragraphs[0].text and not getattr(cell, "_used", False) else cell.add_paragraph()
        cell._used = True  # type: ignore[attr-defined]
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(10)

    put(side, gen.contact_line(cv))
    put(side, h1._text(cv, "skills"), True)
    put(side, ", ".join(cv["skills"]))
    put(side, h1._text(cv, "languages"), True)
    put(side, ", ".join(cv["languages"]))
    put(side, h1._text(cv, "education"), True)
    for ed in cv["education"]:
        put(side, f"{ed['degree']} – {ed['institution']} ({ed['dates']})")
    if cv.get("certifications"):
        put(side, h1._text(cv, "certifications"), True)
        for c in cv["certifications"]:
            put(side, f"{c['name']} — {c['issuer']} ({c['date']})")

    put(main, h1._text(cv, "summary"), True)
    put(main, cv["summary"])
    put(main, h1._text(cv, "experience"), True)
    for e in cv["experience"]:
        put(main, f"{e['title']} — {e['org']}, {e['location']} ({e['dates']})", True)
        for b in e["bullets"]:
            put(main, f"• {b}")
    doc.save(path)


def render_pdf_inline(cv: dict, path: Path) -> None:
    """Bullets under a single header line 'Title at Org, City (dates)'."""
    at = "at" if cv["lang"] == "en" else "di"
    story = [
        Paragraph(cv["identity"]["name"].upper(), gen.ST["name"]),
        Paragraph(cv["headline"], gen.ST["headline"]),
        Spacer(1, 3),
        Paragraph(gen.contact_line(cv), gen.ST["contact"]),
        Paragraph(h1._text(cv, "summary"), gen.ST["h"]),
        Paragraph(cv["summary"], gen.ST["body"]),
        Paragraph(h1._text(cv, "experience"), gen.ST["h"]),
    ]
    for e in cv["experience"]:
        story.append(Paragraph(f"{e['title']} {at} {e['org']}, {e['location']} ({e['dates']})", gen.ST["role"]))
        story.append(gen.bullets(e["bullets"]))
        story.append(Spacer(1, 4))
    for key in ("projects", "organizations"):
        entries = cv.get(key, [])
        if entries:
            story.append(Paragraph(h1._text(cv, key), gen.ST["h"]))
            for e in entries:
                story.append(Paragraph(f"{e['title']} ({e['dates']})", gen.ST["role"]))
                story.append(gen.bullets(e["bullets"]))
    story.append(Paragraph(h1._text(cv, "education"), gen.ST["h"]))
    for ed in cv["education"]:
        story.append(Paragraph(f"<b>{ed['degree']}</b>, {ed['institution']} ({ed['dates']})", gen.ST["body"]))
    if cv.get("certifications"):
        story.append(Paragraph(h1._text(cv, "certifications"), gen.ST["h"]))
        for c in cv["certifications"]:
            story.append(Paragraph(f"{c['name']} — {c['issuer']} ({c['date']})", gen.ST["body"]))
    story += [Paragraph(h1._text(cv, "skills"), gen.ST["h"]), Paragraph(", ".join(cv["skills"]), gen.ST["body"])]
    story += [Paragraph(h1._text(cv, "languages"), gen.ST["h"]), Paragraph(", ".join(cv["languages"]), gen.ST["body"])]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm).build(story)


RENDER = {
    **h1.RENDER,
    "docx_table2": render_docx_table2,
    "docx_dates_first": render_docx_dates_first,
    "docx_sidebar": render_docx_sidebar,
    "pdf_inline": render_pdf_inline,
}


def build_labels() -> dict:
    jd = json.loads((GOLDEN / "job_descriptions/JD-01_requirements.approved.json").read_text())
    candidates, scored = [], []
    for cv in CVS:
        states = cv["expected"]["states"]
        candidates.append({
            "cv_id": cv["id"],
            "file": f"cvs/core/{cv['file']}",
            "group": cv["group"],
            "format": cv["format"],
            "language": cv["lang"],
            "relevance_grade": cv["grade"],
            "expected_states": states,
            "expected_backend_months": cv["expected"]["backend_months"],
            "note": cv["expected"]["note"],
        })
        first = {k: v[0] for k, v in states.items()}
        scored.append({
            "candidate_id": cv["id"],
            "result": score(first, jd),
            "strong_count": sum(v == S for v in first.values()),
        })
    return {
        "job_id": "JD-01",
        "as_of": AS_OF,
        "requirement_set": "docs/spec/golden_dataset/job_descriptions/JD-01_requirements.approved.json",
        "authoring": "Labels derived from docs/spec/02_SCORING_RUBRIC_SPEC.md section 4 before the system was run on "
                     "these CVs. Written by the assistant that worked on the system; needs human review.",
        "reference_ranking": [r["candidate_id"] for r in rank(scored)],
        "candidates": candidates,
    }


def readme(labels: dict) -> str:
    rows = []
    for c in labels["candidates"]:
        st = " ".join("/".join(s[0] for s in c["expected_states"][f"R{i}"]) for i in range(1, 11))
        rows.append(f"| {c['cv_id']} | {c['format']} | {c['language']} | {c['relevance_grade']} | `{st}` | {c['note']} |")
    return (
        "# Second hold-out set\n\n"
        "Twelve more fictional CVs (K01 to K12), scored against JD-01. See the docstring of "
        "`scripts/build_holdout2_set.py` for how it differs from the first hold-out.\n\n"
        "**Measure only. Do not tune rules, prompts, thresholds, or settings against this set.** If the system disagrees "
        "with a label, review the case: either may be wrong.\n\n"
        "**The labels are not independent.** They were written from the scoring spec (section 4) before the system was "
        "run on these CVs, but by the assistant that worked on the system. A human recruiter should review them.\n\n"
        "**It is not a random sample.** It was designed after the first hold-out was measured and deliberately includes "
        "cases for behaviors already under suspicion: ACHIEVEMENT with no numeric result or with numbers in an unrelated "
        "domain, and numeric results written as words or only in a project. K03 uses a date-first layout but has a single "
        "job, so it does not exercise the known multi-job parser defect.\n\n"
        "File names are descriptive for readability, but the system infers a candidate name from the file name; "
        "K09 (`platform_head`) shows the effect: its job title was redacted as if it were a name.\n\n"
        "Rebuild the CV files (PDF and DOCX are git-ignored) with `python scripts/build_holdout2_set.py`, then:\n\n"
        "```\npython scripts/eval_agent_path.py --cvs-dir docs/eval/holdout_set_2/cvs/core "
        "--labels docs/eval/holdout_set_2/labels/expected_JD-01.json --output-json out.json\n```\n\n"
        "## Expected states (R1 to R10: S = STRONG, P = PARTIAL, N = NOT_DEMONSTRATED)\n\n"
        "| CV | Layout | Lang | Grade | R1..R10 | Notes |\n|---|---|---|---|---|---|\n" + "\n".join(rows) + "\n"
    )


def main() -> None:
    ap = argparse.ArgumentParser(description="Build the second hold-out set")
    ap.add_argument("--out", type=Path, default=REPO_ROOT / "docs/eval/holdout_set_2")
    args = ap.parse_args()

    core = args.out / "cvs" / "core"
    core.mkdir(parents=True, exist_ok=True)
    for cv in CVS:
        RENDER[cv["format"]](cv, core / cv["file"])
        print("wrote", cv["file"])

    labels = build_labels()
    (args.out / "labels").mkdir(parents=True, exist_ok=True)
    (args.out / "labels" / "expected_JD-01.json").write_text(json.dumps(labels, indent=2, ensure_ascii=False))
    (args.out / "README.md").write_text(readme(labels))
    print(f"\n{len(CVS)} CVs -> {args.out}; reference ranking: {labels['reference_ranking']}")


if __name__ == "__main__":
    main()
